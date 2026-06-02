"""Markdown → DOCX 转换器（服务化版本）

改造自 md2docx.py，支持：
- 动态 BASE_DIR（不硬编码为脚本目录）
- 动态 MD 文件名列表
- convert_task() 接口用于流水线调用
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def add_heading(doc, text, level):
    h = doc.add_heading(text, min(level, 9))
    return h


def add_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        elif part:
            p.add_run(part)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.3)
    text = '\n'.join(lines)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
    shading = run._element.get_or_add_rPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): 'F5F5F5'
    })
    shading.append(shd)
    return p


def add_table(doc, header, rows):
    ncols = len(header)
    nrows = len(rows) + 1
    table = doc.add_table(rows=nrows, cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.bold = True
        run.font.size = Pt(10)
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(qn('w:shd'), {
            qn('w:val'): 'clear',
            qn('w:color'): 'auto',
            qn('w:fill'): 'E8E8E8'
        })
        shading.append(shd)

    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            if c_idx < ncols:
                cell = table.rows[r_idx + 1].cells[c_idx]
                add_paragraph(cell, cell_text)

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)

    return table


def add_image(doc, alt_text, image_path, base_dir):
    """添加图片，base_dir 为参数传入而非硬编码"""
    full_path = os.path.join(base_dir, image_path)
    if os.path.exists(full_path):
        from PIL import Image
        img = Image.open(full_path)
        w, h = img.size
        max_width = Inches(6.5)
        ratio = min(max_width / Emu(w * 9525), Inches(4.5) / Emu(h * 9525))
        width = int(w * 9525 * ratio)
        height = int(h * 9525 * ratio)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(full_path, width=Emu(width), height=Emu(height))
        cap = doc.add_paragraph(alt_text)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].italic = True
    else:
        p = doc.add_paragraph()
        run = p.add_run(f'[图片缺失: {image_path}]')
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        run.italic = True


def add_list_item(doc, text, ordered=False, level=0):
    style = 'List Number' if ordered else 'List Bullet'
    indent = Inches(0.25 * level)
    p = add_paragraph(doc, text, style=style)
    p.paragraph_format.left_indent = indent
    return p


def parse_md_to_docx(md_file, docx_file, base_dir=None):
    """解析 Markdown 文件并生成 DOCX，base_dir 参数化"""
    if base_dir is None:
        base_dir = os.path.dirname(os.path.abspath(md_file))

    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        if not line.strip():
            i += 1
            continue

        # 图片行 — 传入 base_dir
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            alt_text = img_match.group(1)
            image_path = img_match.group(2)
            add_image(doc, alt_text, image_path, base_dir)
            i += 1
            continue

        heading_match = re.match(r'^#{1,6}\s+(.*)', line)
        if heading_match:
            level = len(re.match(r'^#{1,6}', line).group())
            text = heading_match.group(1)
            add_heading(doc, text, level)
            i += 1
            continue

        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].rstrip('\n').startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            if i < len(lines):
                i += 1
            add_code_block(doc, code_lines)
            continue

        if re.match(r'^\|.*\|$', line):
            table_lines = []
            while i < len(lines) and re.match(r'^\|.*\|$', lines[i].rstrip('\n')):
                table_lines.append(lines[i].rstrip('\n'))
                i += 1
            header = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            data_start = 1
            if len(table_lines) > 1 and all(
                re.match(r'^[\s:-]+$', c.strip()) for c in table_lines[1].split('|')[1:-1]
            ):
                data_start = 2
            rows = []
            for tl in table_lines[data_start:]:
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                rows.append(cells)
            add_table(doc, header, rows)
            continue

        ol_match = re.match(r'^(\s*)(\d+)\.\s+(.*)', line)
        if ol_match:
            indent_level = len(ol_match.group(1)) // 2
            text = ol_match.group(3)
            add_list_item(doc, text, ordered=True, level=indent_level)
            i += 1
            continue

        ul_match = re.match(r'^(\s*)([-*+])\s+(.*)', line)
        if ul_match:
            indent_level = len(ul_match.group(1)) // 2
            text = ul_match.group(3)
            add_list_item(doc, text, ordered=False, level=indent_level)
            i += 1
            continue

        if line.startswith('>'):
            quote_text = line.lstrip('> ').strip()
            while i + 1 < len(lines) and lines[i + 1].rstrip('\n').startswith('>'):
                i += 1
                quote_text += '\n' + lines[i].rstrip('\n').lstrip('> ').strip()
            p = add_paragraph(doc, quote_text)
            p.paragraph_format.left_indent = Inches(0.5)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # 普通段落
        para_text = line
        while i + 1 < len(lines):
            next_line = lines[i + 1].rstrip('\n')
            if not next_line.strip():
                break
            if re.match(r'^#{1,6}', next_line):
                break
            if next_line.startswith('```'):
                break
            if re.match(r'^\|.*\|$', next_line):
                break
            if re.match(r'^\s*(\d+\.|[-*+])\s', next_line):
                break
            if next_line.startswith('>'):
                break
            if re.match(r'!\[.*\]\(.*\)', next_line):
                break
            para_text += ' ' + next_line
            i += 1
        add_paragraph(doc, para_text)
        i += 1

    doc.save(docx_file)
    print(f'已生成: {docx_file}')


def convert_task(task_dir: str, md_files: list[str] = None):
    """流水线调用接口 — 转换指定目录下的 MD 文件为 DOCX

    Args:
        task_dir: 文档输出目录（MD 文件和 images 子目录所在位置）
        md_files: MD 文件名列表，如为 None 则自动扫描目录中所有 .md 文件
    """
    if md_files is None:
        md_files = [f for f in os.listdir(task_dir) if f.endswith('.md')]

    for md_file in md_files:
        md_path = os.path.join(task_dir, md_file)
        docx_name = os.path.splitext(md_file)[0] + '.docx'
        docx_path = os.path.join(task_dir, docx_name)
        if os.path.exists(md_path):
            parse_md_to_docx(md_path, docx_path, base_dir=task_dir)

    return [os.path.splitext(f)[0] + '.docx' for f in md_files if os.path.exists(os.path.join(task_dir, f))]